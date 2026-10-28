"""Exportaciones Excel (openpyxl). Solo da formato a filas ya autorizadas; no consulta la base de datos."""
import io
from datetime import date, datetime
from decimal import Decimal
from uuid import UUID

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

TEMPLATE_VERSION = "excel/1.0"
_HEAD = PatternFill("solid", fgColor="0B1F3A")


def _cell(value):
    if isinstance(value, UUID):
        return str(value)
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, (list, tuple)):
        return ", ".join(str(v) for v in value)
    if isinstance(value, dict):
        return "; ".join(f"{k}: {v}" for k, v in value.items())
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@"):
        # Evita inyección de fórmulas al abrir el archivo en Excel.
        return "'" + value
    return value


def workbook(sheets: list[tuple[str, list[str], list[list]]], meta: dict) -> bytes:
    wb = Workbook()
    info = wb.active
    info.title = "Información"
    for row in (("8D PRO", meta["organization"]), ("Documento", meta["title"]), ("Folio", meta["folio"]),
                ("Revisión", meta["revision"]), ("Emitido", meta["issued_at"]), ("Emitido por", meta["issued_by"]),
                ("Plantilla", TEMPLATE_VERSION), ("Filtros", meta.get("filters", "—")),
                ("Nota", "Exportación de datos autorizados para el usuario emisor. La huella SHA-256 queda registrada en 8D PRO.")):
        info.append(row)
    info["A1"].font = Font(bold=True, size=13, color="0B1F3A")
    info.column_dimensions["A"].width, info.column_dimensions["B"].width = 18, 90
    for title, header, rows in sheets:
        ws = wb.create_sheet(title[:31])
        ws.append(header)
        for c in ws[1]:
            c.font, c.fill, c.alignment = Font(bold=True, color="FFFFFF"), _HEAD, Alignment(vertical="center")
        for r in rows:
            ws.append([_cell(v) for v in r])
        ws.freeze_panes = "A2"
        if rows:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}{len(rows) + 1}"
        for i, h in enumerate(header, start=1):
            width = max([len(str(h))] + [len(str(_cell(r[i - 1]) or "")) for r in rows[:200]])
            ws.column_dimensions[get_column_letter(i)].width = min(max(10, width + 2), 60)
            for cell in ws[get_column_letter(i)][1:]:
                if isinstance(cell.value, datetime):
                    cell.number_format = "yyyy-mm-dd hh:mm"
                elif isinstance(cell.value, date):
                    cell.number_format = "yyyy-mm-dd"
    wb.properties.creator = "8D PRO"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
