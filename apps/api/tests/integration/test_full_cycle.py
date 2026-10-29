"""Flujo crítico completo D0-D8: acciones, eficacia, lección aprendida, documentos, Excel, bitácora y versiones."""
import io
from pathlib import Path

import psycopg
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

from app import storage

from .conftest import PASSWORD, png


def ok(res, status=200):
    assert res.status_code == status, res.text
    return res.json()


def flow(case, code):
    return next(d for d in case["workflow"] if d["code"] == code)


def pdf_text(content: bytes) -> str:
    return "\n".join(page.extract_text() for page in PdfReader(io.BytesIO(content)).pages)


@pytest.fixture
def actors(login):
    return {"ing": login("ing@t.local"), "ing2": login("ing2@t.local"), "mgr": login("gerente@t.local"),
            "admin": login("admin@t.local"), "viewer": login("lector@t.local"), "other": login("otra@x.local")}


def approve(a, cid, code):
    ok(a["ing"].post(f"/cases/{cid}/disciplines/{code}/submit"))
    return ok(a["mgr"].post(f"/cases/{cid}/disciplines/{code}/approve", json={"comment": "OK"}))


def through_d4(a, plant, users) -> tuple[str, dict]:
    ing, mgr = a["ing"], a["mgr"]
    case = ok(ing.post("/cases", json={"plant_id": str(plant), "title": "Rebaba en carcasa 123", "process_name": "Maquinado CNC",
                                       "source_type": "internal", "part_number": "123-A", "product_name": "Carcasa",
                                       "customer_name": "Cliente A", "target_close_date": "2026-12-15"}), 201)
    cid = case["id"]
    ok(ing.patch(f"/cases/{cid}/disciplines/D0", json={"content": {"emergency_response": "Bloqueo", "justification": "Recurrente"}}))
    approve(a, cid, "D0")
    ok(ing.put(f"/cases/{cid}/team", json={"members": [
        {"user_id": str(users["ing@t.local"]), "team_role": "leader"}, {"user_id": str(users["ing2@t.local"])},
        {"user_id": str(users["gerente@t.local"]), "team_role": "champion"}]}))
    case = approve(a, cid, "D1")
    ok(ing.patch(f"/cases/{cid}", json={"version": case["version"], "problem_description": "Rebaba de 0.5 mm en el borde exterior",
                                        "defect_date": "2026-10-01", "detected_quantity": 40, "affected_quantity": 1200}))
    ok(ing.patch(f"/cases/{cid}/disciplines/D2", json={"content": {"what": "Rebaba", "where": "Línea 3", "when": "T2",
                                                                    "how_detected": "Visual"}}))
    approve(a, cid, "D2")
    act = ok(ing.post(f"/cases/{cid}/containment", json={"description": "Inspección 100%"}), 201)[0]
    ok(ing.patch(f"/containment/{act['id']}", json={"status": "implemented", "evidence_notes": "R-12"}))
    ok(ing.patch(f"/containment/{act['id']}", json={"status": "verified"}))
    alert = ok(ing.post(f"/cases/{cid}/alerts", json={"failure_type": "internal"}), 201)
    ok(ing.patch(f"/alerts/{alert['id']}", json={"where_happened": "Línea 3", "who_is_affected": "Ensamble",
                                                 "how_detected": "Visual", "containment_summary": "Bloqueo"}))
    photo = ok(ing.post(f"/cases/{cid}/attachments", data={"purpose": "photo"}, files={"file": ("nok.png", png(), "image/png")}), 201)
    ok(ing.put(f"/alerts/{alert['id']}/photos", json={"photos": [{"attachment_id": photo["id"], "kind": "nok"}]}))
    ok(mgr.post(f"/alerts/{alert['id']}/issue"))
    ok(ing.post(f"/cases/{cid}/containment-record/start", json={}))
    ok(ing.post(f"/cases/{cid}/containment-lots", json={"lot_number": "L-0915", "quantity_inspected": 1200, "quantity_rejected": 41}), 201)
    ok(ing.post(f"/cases/{cid}/containment-record/end", json={"result": "effective", "result_notes": "Lotes segregados y revisados"}))
    approve(a, cid, "D3")
    cats = {c["code"]: c["id"] for c in ok(ing.get("/ishikawa/categories"))}
    causes = ok(ing.post(f"/cases/{cid}/causes", json={"category_id": cats["MACHINE"], "potential_cause": "Desgaste de herramienta"}), 201)
    causes = ok(ing.post(f"/cases/{cid}/causes", json={"category_id": cats["MANPOWER"], "potential_cause": "Operador nuevo"}), 201)
    root, other = causes[0], causes[1]
    evidence = {"verification_method": "Reproducción", "evidence_summary": "10/10", "rationale": "Correlación directa"}
    ok(ing.post(f"/causes/{root['id']}/validations", json={"decision": "root_cause", **evidence}), 201)
    ok(ing.post(f"/causes/{other['id']}/validations", json={"decision": "not_root_cause", **evidence}), 201)
    ok(ing.patch(f"/cases/{cid}/disciplines/D4", json={"content": {"occurrence_cause": "Vida útil excedida",
                                                                    "escape_point": "Sin criterio de rebaba"}}))
    approve(a, cid, "D4")
    return cid, {"root": root["id"], "other": other["id"]}


def test_complete_8d_cycle(actors, plant, users, sql):
    a = actors
    ing, ing2, mgr, viewer, other = a["ing"], a["ing2"], a["mgr"], a["viewer"], a["other"]
    cid, causes = through_d4(a, plant, users)

    # ---- D5: acciones correctivas vinculadas a la causa raíz + criterio de eficacia aprobado ----
    wrong = ing.post(f"/cases/{cid}/actions", json={"action_type": "corrective", "cause_id": causes["other"], "description": "Capacitar"})
    assert wrong.status_code == 409 and wrong.json()["code"] == "cause_not_root"
    assert ing.post(f"/cases/{cid}/actions", json={"action_type": "corrective", "description": "Sin causa"}).status_code == 422
    assert viewer.post(f"/cases/{cid}/actions", json={"action_type": "corrective", "cause_id": causes["root"],
                                                      "description": "x" * 5}).status_code == 403
    actions = ok(ing.post(f"/cases/{cid}/actions", json={
        "action_type": "corrective", "cause_id": causes["root"], "description": "Control de vida útil de herramienta en CNC",
        "responsible_user_id": str(users["ing2@t.local"]), "due_date": "2026-11-15"}), 201)
    action_id = actions[0]["id"]
    eff = ok(ing.post(f"/cases/{cid}/criteria", json={
        "description": "Rebaba ≤ 100 ppm durante 30 días con producción", "metric_name": "Piezas con rebaba",
        "defect_unit": "piezas", "denominator_type": "units_produced", "threshold_type": "max_rate_ppm",
        "threshold_value": 100, "min_period_days": 30, "min_denominator": 1000}), 201)
    criterion = eff["criteria"][0]["id"]
    assert ing.post(f"/criteria/{criterion}/approve").status_code == 403
    ok(mgr.post(f"/criteria/{criterion}/approve"))
    with pytest.raises(psycopg.errors.RaiseException):
        sql("update effectiveness_criteria set threshold_value = 999 where id = %s", criterion)
    blocked = ing.post(f"/cases/{cid}/disciplines/D5/submit")
    assert blocked.status_code == 409 and "Plan para verificar" in " ".join(blocked.json()["details"])
    ok(ing.patch(f"/cases/{cid}/disciplines/D5", json={"content": {"verification_plan": "Prueba piloto 2 semanas"}}))
    approve(a, cid, "D5")
    locked = ing.post(f"/cases/{cid}/actions", json={"action_type": "corrective", "cause_id": causes["root"], "description": "Otra"})
    assert locked.json()["code"] == "discipline_locked"

    # ---- D6: implementación con evidencia, aprobación por otra persona y eficacia ----
    early = ing.post(f"/actions/{action_id}/checks", json={"criterion_id": criterion, "period_start": "2026-09-01",
                                                          "period_end": "2026-09-30"})
    assert early.json()["code"] == "not_implemented"
    ok(ing2.post(f"/actions/{action_id}/start"))
    assert ing2.post(f"/actions/{action_id}/submit", json={}).json()["code"] == "evidence_required"
    ok(ing2.post(f"/actions/{action_id}/submit", json={"implementation_evidence": "OT-551: contador de golpes instalado"}))
    assert ing2.post(f"/actions/{action_id}/approve", json={}).status_code == 403
    assert mgr.post(f"/actions/{action_id}/reject", json={}).status_code == 422
    ok(mgr.post(f"/actions/{action_id}/approve", json={"comment": "Verificado en piso"}))
    ok(ing2.post(f"/cases/{cid}/attachments", data={"purpose": "evidence", "action_id": action_id},
                 files={"file": ("ot.pdf", b"%PDF-1.4 OT-551", "application/pdf")}), 201)

    # Datos faltantes: nunca «eficaz».
    eff = ok(ing.post(f"/actions/{action_id}/checks", json={
        "criterion_id": criterion, "period_start": "2026-08-01", "period_end": "2026-09-30", "denominator_quantity": 25000,
        "data_notes": "Falta el conteo de defectos de agosto"}), 201)
    missing = eff["checks"][0]
    assert missing["computed_result"] == "insufficient_data" and missing["allowed_decisions"] == ["insufficient_data"]
    bad = ing.post(f"/checks/{missing['id']}/decide", json={"result": "effective", "rationale": "Parece bien"})
    assert bad.status_code == 409 and bad.json()["code"] == "result_not_supported"
    ok(ing.post(f"/checks/{missing['id']}/decide", json={"result": "insufficient_data", "rationale": "Sin conteo de defectos"}))
    assert ing.post(f"/actions/{action_id}/checks", json={"criterion_id": criterion, "period_start": "2026-09-01",
                                                         "period_end": "2026-09-30", "denominator_quantity": 10,
                                                         "defect_quantity": 11}).status_code == 422
    eff = ok(ing.post(f"/actions/{action_id}/checks", json={
        "criterion_id": criterion, "period_start": "2026-08-01", "period_end": "2026-09-30", "denominator_quantity": 25000,
        "defect_quantity": 1, "days_without_production": 4}), 201)
    good = next(c for c in eff["checks"] if c["result"] == "pending")
    assert good["computed_result"] == "effective" and good["evaluation"]["value"] == pytest.approx(40)
    ok(ing.post(f"/cases/{cid}/attachments", data={"purpose": "evidence", "effectiveness_check_id": good["id"]},
                files={"file": ("datos.pdf", b"%PDF-1.4 datos", "application/pdf")}), 201)
    ok(ing.post(f"/checks/{good['id']}/decide", json={"result": "effective", "rationale": "40 ppm < 100 ppm en 57 días con producción"}))
    assert ing.post(f"/checks/{good['id']}/decide", json={"result": "ineffective", "rationale": "cambio"}).json()["code"] == "already_decided"
    with pytest.raises(psycopg.errors.RaiseException):
        sql("update effectiveness_checks set result = 'ineffective' where id = %s", good["id"])
    assert next(x for x in ok(ing.get(f"/cases/{cid}/actions")) if x["id"] == action_id)["status"] == "verified"
    ok(ing.patch(f"/cases/{cid}/disciplines/D6", json={"content": {"validation_summary": "Eficaz según criterio aprobado"}}))
    approve(a, cid, "D6")

    # ---- D7: preventiva, actualización documental y despliegue horizontal ----
    prev = ok(ing.post(f"/cases/{cid}/actions", json={"action_type": "preventive", "description": "Estandarizar vida útil en todas las CNC",
                                                      "responsible_user_id": str(users["ing@t.local"]), "due_date": "2026-12-01"}), 201)
    prev_id = next(x["id"] for x in prev if x["action_type"] == "preventive")
    ok(ing.post(f"/actions/{prev_id}/submit", json={"evidence_exception": "Cambio de software sin registro físico; validado en sitio"}))
    assert ing.post(f"/actions/{prev_id}/approve", json={}).status_code == 403
    self_approve = a["admin"].post(f"/actions/{prev_id}/approve", json={})
    assert self_approve.status_code == 200  # admin no fue quien la envió
    upd = ok(ing.post(f"/cases/{cid}/system-updates", json={"update_type": "pfmea", "document_reference": "PFMEA-CNC-12"}), 201)
    no_ev = ing.patch(f"/system-updates/{upd[0]['id']}", json={"status": "completed"})
    assert no_ev.json()["code"] == "evidence_required"
    ok(ing.patch(f"/system-updates/{upd[0]['id']}", json={"status": "completed", "completed_on": "2026-10-20",
                                                          "evidence_notes": "PFMEA-CNC-12 rev D, RPN 120 → 40"}))
    ok(ing.patch(f"/cases/{cid}/disciplines/D7", json={"content": {"horizontal_deployment": "Aplicado a CNC 1-6 de P1 y P2"}}))
    approve(a, cid, "D7")

    # ---- D8 y lección aprendida ----
    ok(ing.patch(f"/cases/{cid}/disciplines/D8", json={"content": {"closure_summary": "Problema resuelto",
                                                                    "team_recognition": "Reconocimiento al equipo de CNC"}}))
    d8 = ing.post(f"/cases/{cid}/disciplines/D8/submit")
    assert d8.status_code == 409 and "Lección aprendida aprobada o publicada" in d8.json()["details"]
    lesson = ok(ing.post(f"/cases/{cid}/lessons"), 201)
    lid = lesson["id"]
    assert lesson["folio"].startswith("LA-") and "Desgaste de herramienta" in lesson["verified_root_cause"]
    assert "PFMEA-CNC-12" in lesson["documents_updated"] and "eficaz" in lesson["effectiveness_summary"]
    assert viewer.get(f"/lessons/{lid}").status_code == 404  # borrador no visible para consulta
    assert ing.post(f"/lessons/{lid}/submit").json()["code"] == "requirements_pending"
    ok(ing.patch(f"/lessons/{lid}", json={"before_summary": "1.2 % de rebaba", "after_summary": "40 ppm",
                                          "situations": ["Herramental con vida útil"], "language_code": "es"}))
    no_photos = ing.post(f"/lessons/{lid}/submit")
    assert no_photos.status_code == 409 and {"Foto del antes", "Foto del después"} <= set(no_photos.json()["details"])
    assert ing.post(f"/lessons/{lid}/photos/upload", data={"kind": "before"},
                    files={"file": ("x.pdf", b"%PDF-1.4", "application/pdf")}).status_code == 422
    ok(ing.post(f"/lessons/{lid}/photos/upload", data={"kind": "before", "caption": "Rebaba en el borde"},
                files={"file": ("antes.png", png((200, 40, 40)), "image/png")}), 201)
    lesson = ok(ing.post(f"/lessons/{lid}/photos/upload", data={"kind": "after", "caption": "Borde limpio"},
                         files={"file": ("despues.png", png((40, 160, 60)), "image/png")}), 201)
    assert [p["kind"] for p in lesson["photos"]] == ["before", "after"] and lesson["missing_to_submit"] == []
    ok(ing.post(f"/lessons/{lid}/submit"))
    assert ing.put(f"/lessons/{lid}/photos", json={"photos": []}).json()["code"] == "lesson_locked"
    with pytest.raises(psycopg.errors.RaiseException):
        sql("delete from lesson_photos where lesson_id = %s", lid)
    assert ing.patch(f"/lessons/{lid}", json={"title": "Cambio tras enviar"}).json()["code"] == "lesson_locked"
    assert ing.post(f"/lessons/{lid}/review", json={"decision": "approve"}).status_code == 403
    assert mgr.post(f"/lessons/{lid}/review", json={"decision": "reject"}).status_code == 422
    ok(mgr.post(f"/lessons/{lid}/review", json={"decision": "approve", "comment": "Clara y verificable"}))
    case = approve(a, cid, "D8")
    assert case["status"] == "closed" and case["closed_at"]
    assert ing.patch(f"/cases/{cid}/disciplines/D8", json={"content": {"notes": "x"}}).json()["code"] == "case_closed"

    assert ing.post(f"/lessons/{lid}/publish").status_code == 403
    lesson = ok(mgr.post(f"/lessons/{lid}/publish"))
    assert lesson["status"] == "published" and viewer.get(f"/lessons/{lid}").status_code == 200
    assert [x["id"] for x in ok(viewer.get("/lessons", params={"process": "maquinado"}))] == [lid]
    with pytest.raises(psycopg.errors.RaiseException):
        sql("update lessons_learned set title = 'alterado' where id = %s", lid)

    # Reutilización en otro caso: registra evaluación; no modifica lección ni expediente original.
    target = ok(ing.post("/cases", json={"plant_id": str(plant), "title": "Rebaba en tapa"}), 201)
    before = sql("select updated_at from lessons_learned where id = %s", lid)[0]["updated_at"]
    assert ing.post(f"/lessons/{lid}/apply", json={"target_case_id": cid, "evaluation": "mismo caso", "outcome": "applied"}).status_code == 422
    lesson = ok(ing.post(f"/lessons/{lid}/apply", json={"target_case_id": target["id"], "outcome": "partially_applied",
                                                        "evaluation": "Se revisó el contador de golpes en la prensa 4"}), 201)
    assert lesson["applications"][0]["target_folio"] == target["folio"]
    assert ing.post(f"/lessons/{lid}/apply", json={"target_case_id": target["id"], "outcome": "applied",
                                                   "evaluation": "otra vez"}).json()["code"] == "duplicate"
    assert sql("select updated_at from lessons_learned where id = %s", lid)[0]["updated_at"] == before
    assert ok(ing.get(f"/cases/{target['id']}/lessons"))["applied"][0]["lesson_id"] == lid

    # Nueva revisión: la publicada sigue vigente hasta publicar la nueva, que la sustituye.
    rev2 = ok(ing.post(f"/lessons/{lid}/revise", json={"reason": "Agregar resultado de P2"}), 201)
    assert rev2["revision"] == 2 and rev2["status"] == "draft" and len(rev2["photos"]) == 2  # las fotos se copian
    assert ing.post(f"/lessons/{rev2['id']}/apply", json={"target_case_id": target["id"], "outcome": "applied",
                                                          "evaluation": "borrador"}).json()["code"] == "not_published"
    ok(ing.patch(f"/lessons/{rev2['id']}", json={"after_summary": "40 ppm en P1 y 35 ppm en P2"}))
    ok(ing.post(f"/lessons/{rev2['id']}/submit"))
    ok(mgr.post(f"/lessons/{rev2['id']}/review", json={"decision": "approve"}))
    rev2 = ok(mgr.post(f"/lessons/{rev2['id']}/publish"))
    assert {r["revision"]: r["status"] for r in rev2["revisions"]} == {2: "published", 1: "superseded"}

    # ---- Documentos: 8D completo, avance, auditoría, lección, Excel; versiones e integridad ----
    preview = ing.get(f"/cases/{cid}/documents/8d_full/preview")
    assert preview.headers["content-type"] == "application/pdf" and "VISTA PREVIA" in pdf_text(preview.content)
    assert not sql("select 1 from documents where case_id = %s and document_type = '8d_full'", cid)
    full1 = ok(ing.post(f"/cases/{cid}/documents/8d_full"), 201)
    text = pdf_text(ing.get(f"/documents/{full1['id']}/content").content)
    for expected in ("Expediente 8D completo", "D8 — Cierre y reconocimiento", "Control de vida útil", "PFMEA-CNC-12",
                     "Eficaz", "LA-", "Página 1 de", "Cerrado"):
        assert expected in text, expected
    full2 = ok(ing.post(f"/cases/{cid}/documents/8d_full"), 201)
    assert (full1["revision"], full2["revision"]) == (1, 2)
    status = {r["id"]: r["status"] for r in sql("select id::text, status from documents where case_id = %s", cid)}
    assert status[full1["id"]] == "superseded" and status[full2["id"]] == "issued"
    assert ing.get(f"/documents/{full1['id']}/content").status_code == 200  # la emisión anterior se conserva
    progress = ok(ing.post(f"/cases/{cid}/documents/8d_progress"), 201)
    assert "Reporte de avance" in pdf_text(ing.get(f"/documents/{progress['id']}/content").content)
    assert ing.post(f"/cases/{cid}/documents/audit_report").status_code == 403
    audit_doc = ok(mgr.post(f"/cases/{cid}/documents/audit_report"), 201)
    assert "Reporte de auditoría" in pdf_text(mgr.get(f"/documents/{audit_doc['id']}/content").content)
    lesson_doc = ok(ing.post(f"/lessons/{rev2['id']}/document"), 201)
    lesson_pdf = PdfReader(io.BytesIO(ing.get(f"/documents/{lesson_doc['id']}/content").content))
    assert "ANTES" in lesson_pdf.pages[0].extract_text() + lesson_pdf.pages[-1].extract_text()
    assert sum(len(pg.images) for pg in lesson_pdf.pages) >= 2
    draft_lesson = ok(ing.post(f"/lessons/{rev2['id']}/revise", json={"reason": "Tercera revisión"}), 201)
    assert ing.post(f"/lessons/{draft_lesson['id']}/document").json()["code"] == "lesson_not_approved"

    xlsx = ok(ing.post(f"/cases/{cid}/export"), 201)
    wb = load_workbook(io.BytesIO(ing.get(f"/documents/{xlsx['id']}/content").content))
    assert {"Información", "Caso", "Acciones", "Eficacia", "Bitácora"} <= set(wb.sheetnames)
    assert wb["Acciones"].max_row >= 3
    listing = ok(ing.post("/exports/cases", json={"q": "Rebaba"}), 201)
    wb = load_workbook(io.BytesIO(ing.get(f"/documents/{listing['id']}/content").content))
    assert wb["Casos"].max_row == 3  # encabezado + 2 casos
    assert viewer.post("/exports/cases", json={}).status_code == 403
    assert ing.post("/exports/audit", json={}).status_code == 403
    assert ok(mgr.post("/exports/audit", json={}), 201)["document_type"] == "excel_export"
    # Exportación de organización: no la descarga otro usuario sin auditoría ni otra organización.
    assert viewer.get(f"/documents/{listing['id']}/content").status_code == 404
    assert other.get(f"/documents/{full2['id']}/content").status_code == 404
    assert other.get(f"/documents/{full2['id']}/verify").status_code == 404

    assert ok(ing.get(f"/documents/{full2['id']}/verify"))["ok"] is True
    key = sql("select storage_key from documents where id = %s", full2["id"])[0]["storage_key"]
    Path(storage.root() / key).write_bytes(b"%PDF-alterado")
    tampered = ok(ing.get(f"/documents/{full2['id']}/verify"))
    assert tampered["ok"] is False and "alterado" in tampered["detail"]

    # ---- Bitácora ----
    events = {e["event_type"] for e in ok(ing.get(f"/cases/{cid}/audit"))}
    assert {"action.created", "action.submitted", "action.approved", "criterion.approved", "effectiveness.decided",
            "system_update.updated", "lesson.created", "lesson.approved", "lesson.published", "document.issued",
            "document.downloaded", "integrity.verified", "discipline.approved"} <= events
    txids = sql("select metadata->>'txid' as t, request_id from audit_events where case_id = %s and event_type = 'document.issued'", cid)
    assert all(r["t"] and r["request_id"] for r in txids)

    # ---- Indicadores ----
    k = ok(ing.get("/reports/kpis"))
    assert k["closed_cases"] == 1 and k["open_cases"] == 1
    assert k["resolution"] == {"n": 1, "sufficient": False}  # sin cálculo con menos de 3 casos
    assert k["effectiveness"]["effective"] == 1 and k["effectiveness"]["effective_pct"] is None
    assert {d["id"] for d in k["definitions"]} >= {"open_cases", "resolution_days", "effectiveness"}

    # ---- Reapertura del caso cerrado solo por D8 ----
    case = ok(mgr.post(f"/cases/{cid}/disciplines/D8/reopen", json={"comment": "Reclamo de seguimiento"}))
    assert case["status"] == "in_progress" and case["closed_at"] is None


def test_document_generation_failure_registers_nothing(actors, plant, sql, monkeypatch):
    from app import pdf_reports
    ing = actors["ing"]
    case = ok(ing.post("/cases", json={"plant_id": str(plant), "title": "x"}), 201)

    def boom(*_):
        raise RuntimeError("plantilla dañada")
    monkeypatch.setattr(pdf_reports, "case_report", boom)
    res = ing.post(f"/cases/{case['id']}/documents/8d_progress")
    assert res.status_code == 500 and res.json()["code"] == "document_failed"
    assert not sql("select 1 from documents where case_id = %s", case["id"])
    assert not sql("select 1 from audit_events where case_id = %s and event_type = 'document.issued'", case["id"])


def test_login_lockout_and_security_events(login, sql):
    admin = login("admin@t.local")
    ok(admin.post("/admin/users", json={"email": "bloqueo@t.local", "display_name": "Bloqueo",
                                        "password": "Contrasena-Larga-9", "roles": ["viewer"]}), 201)
    from fastapi.testclient import TestClient
    c = TestClient(admin.app, headers={"x-d8-client": "web"})
    for _ in range(5):
        assert c.post("/auth/login", json={"email": "bloqueo@t.local", "password": "mala"}).status_code == 401
    locked = c.post("/auth/login", json={"email": "bloqueo@t.local", "password": "Contrasena-Larga-9"})
    assert locked.status_code == 429 and locked.json()["code"] == "too_many_attempts"
    events = [r["event_type"] for r in sql("""select e.event_type from audit_events e join users u on u.id = e.actor_user_id
                                              where u.email = 'bloqueo@t.local' order by e.id""")]
    assert events.count("auth.login_failed") == 5 and "auth.locked" in events
    sql("delete from login_attempts where email = 'bloqueo@t.local'")
    assert c.post("/auth/login", json={"email": "bloqueo@t.local", "password": "Contrasena-Larga-9"}).status_code == 200
    assert c.post("/cases", json={"plant_id": "00000000-0000-0000-0000-000000000000", "title": "x"}).status_code == 403
    assert sql("""select 1 from audit_events e join users u on u.id = e.actor_user_id
                  where u.email = 'bloqueo@t.local' and e.event_type = 'security.forbidden'""")


def test_actions_respect_organization_isolation(actors, plant, users, sql):
    cid, causes = through_d4(actors, plant, users)
    other = actors["other"]
    actions = ok(actors["ing"].post(f"/cases/{cid}/actions", json={"action_type": "corrective", "cause_id": causes["root"],
                                                                    "description": "Acción privada"}), 201)
    aid = actions[0]["id"]
    ok(actors["ing"].patch(f"/actions/{aid}", json={"due_date": "2026-01-10"}))
    dash = ok(actors["ing"].get("/dashboard"))
    assert dash["attention"]["actions_overdue"] == 1 and dash["attention_items"][0]["overdue"] is True
    assert ok(other.get("/dashboard"))["attention"]["actions_overdue"] == 0
    assert other.get(f"/cases/{cid}/actions").status_code == 404
    assert other.post(f"/actions/{aid}/start").status_code == 404
    assert other.get(f"/cases/{cid}/effectiveness").status_code == 404
    assert other.post(f"/cases/{cid}/lessons").status_code == 404
    assert other.get("/lessons").json() == []
    assert other.get("/reports/kpis").json()["open_cases"] == 0
    # El mismo correo con contraseña ajena no entra.
    assert other.post("/auth/login", json={"email": "ing@t.local", "password": PASSWORD + "x"}).status_code == 401
