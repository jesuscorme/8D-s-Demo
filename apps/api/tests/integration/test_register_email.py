"""8D internos/externos, incidencias del Ishikawa, concentrado con hitos, tablero y correo."""
import io

from openpyxl import load_workbook

from .test_full_cycle import ok, through_d4


def test_internal_external_and_register_milestones(login, plant, users, sql):
    actors = {"ing": login("ing@t.local"), "mgr": login("gerente@t.local")}
    ing = actors["ing"]
    cid, _ = through_d4(actors, plant, users)  # interno, D0-D4 aprobadas hoy
    ext = ok(ing.post("/cases", json={"plant_id": str(plant), "title": "Reclamo externo", "case_type": "external",
                                      "market_segment": "OEM", "main_defect": "Rebaba", "report_date": "2026-01-05",
                                      "customer_name": "Cliente Uno"}), 201)
    assert ext["case_type"] == "external" and ext["market_segment"] == "OEM"
    assert ing.post("/cases", json={"plant_id": str(plant), "title": "x", "case_type": "otro"}).status_code == 422
    assert ing.post("/cases", json={"plant_id": str(plant), "title": "x", "report_date": "2099-01-01"}).status_code == 422
    assert [c["id"] for c in ok(ing.get("/cases", params={"case_type": "external"}))] == [ext["id"]]

    reg = ok(ing.get("/reports/register"))
    assert [m["days"] for m in reg["milestones"]] == [1, 2, 7, 14]
    rows = {r["id"]: r for r in reg["rows"]}
    internal = {m["key"]: m["status"] for m in rows[cid]["milestones"]}
    assert internal["containment"] == "on_time" and internal["root_cause"] == "on_time"
    assert rows[cid]["root_cause"] == "Desgaste de herramienta" and rows[cid]["leader"] == "Ingeniera"
    old = {m["key"]: m["status"] for m in rows[ext["id"]]["milestones"]}
    assert set(old.values()) == {"overdue"} and rows[ext["id"]]["overdue"] is True  # reportado en enero, sin avance
    assert [r["id"] for r in ok(ing.get("/reports/register", params={"case_type": "internal"}))["rows"]] == [cid]

    doc = ok(ing.post("/exports/register", json={}), 201)
    wb = load_workbook(io.BytesIO(ing.get(f"/documents/{doc['id']}/content").content))
    header = [c.value for c in wb["Concentrado"][1]]
    assert header[:6] == ["#", "Mes", "Fecha de reporte", "Tipo", "Cliente", "Descripción de la falla"]
    assert wb["Concentrado"].max_row == 3

    dash = ok(ing.get("/dashboard"))
    assert dash["totals"] == {"total": 2, "in_process": 2, "closed": 0, "overdue": 1}
    assert {g["key"] for g in dash["status_groups"]} == {"in_process", "review", "closed"}
    assert {c["label"] for c in dash["by_customer"]} == {"Cliente Uno", "Internos (planta)"}
    assert len(dash["trend"]) == 6 and dash["by_type"] == {"internal": 1, "external": 1}
    assert dash["recent"][0]["folio"] and all("due" in r for r in dash["recent"])


def test_ishikawa_incidences_and_history(login, plant, users):
    actors = {"ing": login("ing@t.local"), "mgr": login("gerente@t.local")}
    ing = actors["ing"]
    first, causes = through_d4(actors, plant, users)  # «Desgaste de herramienta» = causa raíz
    case = ok(ing.post("/cases", json={"plant_id": str(plant), "title": "Rebaba otra vez"}), 201)
    cats = {c["code"]: c["id"] for c in ok(ing.get("/ishikawa/categories"))}
    rows = ok(ing.post(f"/cases/{case['id']}/causes", json={"category_id": cats["MACHINE"], "potential_cause": "desgaste de herramienta",
                                                            "incidence_count": 12}), 201)
    c = rows[0]
    assert c["incidence_count"] == 12 and c["history_cases"] == 1 and c["history_root_cases"] == 1
    assert ing.patch(f"/causes/{c['id']}", json={"incidence_count": -1}).status_code == 422
    evidence = {"verification_method": "m", "evidence_summary": "e", "rationale": "r"}
    ok(ing.post(f"/causes/{c['id']}/validations", json={"decision": "root_cause", **evidence}), 201)
    # Decidida: el texto ya no cambia, pero las incidencias sí.
    assert ing.patch(f"/causes/{c['id']}", json={"potential_cause": "otra"}).json()["code"] == "cause_decided"
    assert ok(ing.patch(f"/causes/{c['id']}", json={"incidence_count": 15}))[0]["incidence_count"] == 15
    hist = next(x for x in ok(ing.get(f"/cases/{first}/causes")) if x["id"] == causes["root"])
    assert hist["history_cases"] == 1 and hist["history_root_cases"] == 1


def test_email_admin_test_failure_and_retry(login, sql, monkeypatch):
    admin, ing = login("admin@t.local"), login("ing@t.local")
    assert ing.get("/admin/notifications").status_code == 403
    data = ok(admin.get("/admin/notifications"))
    assert data["settings"]["transport"] == "file" and {r["email"] for r in data["recipients"]} >= {"logistica@t.local"}
    assert admin.post("/admin/notifications", json={"event": "alert_issued", "email": "no-es-correo"}).status_code == 422
    data = ok(admin.post("/admin/notifications", json={"event": "alert_issued", "email": "Nuevo@T.local"}), 201)
    new = next(r for r in data["recipients"] if r["email"] == "nuevo@t.local")
    assert admin.post("/admin/notifications", json={"event": "alert_issued", "email": "nuevo@t.local"}).status_code == 409
    ok(admin.delete(f"/admin/notifications/{new['id']}"))

    sent = ok(admin.post("/admin/email/test", json={"to": "destino@t.local"}), 201)
    assert sql("select status from email_outbox where id = %s", sent["id"])[0]["status"] == "sent"

    # SMTP inaccesible: el correo queda «failed» con su error y se puede reintentar.
    monkeypatch.setenv("EMAIL_TRANSPORT", "smtp")
    monkeypatch.setenv("SMTP_HOST", "127.0.0.1")
    monkeypatch.setenv("SMTP_PORT", "9")
    failed = ok(admin.post("/admin/email/test", json={"to": "destino@t.local"}), 201)
    row = sql("select status, last_error, transport from email_outbox where id = %s", failed["id"])[0]
    assert row["status"] == "failed" and row["last_error"] and row["transport"] == "smtp"
    monkeypatch.setenv("EMAIL_TRANSPORT", "file")
    ok(admin.post(f"/admin/email/outbox/{failed['id']}/retry"))
    assert sql("select status, attempts from email_outbox where id = %s", failed["id"])[0] == {"status": "sent", "attempts": 2}
    assert admin.post(f"/admin/email/outbox/{failed['id']}/retry").json()["code"] == "not_failed"
    assert any(e["id"] == failed["id"] for e in ok(admin.get("/admin/email/outbox")))
