"""Regla de eficacia (datos faltantes ≠ cero, denominador, días sin producción) y servicios de documentos."""
import io
from datetime import date, datetime

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

from app import excel, pdf_reports
from app.effectiveness_rules import InvalidData, allowed_decisions, evaluate

PPM = {"threshold_type": "max_rate_ppm", "threshold_value": 100, "min_period_days": 30, "min_denominator": 1000,
       "denominator_type": "units_produced"}
ZERO = {**PPM, "threshold_type": "max_defects", "threshold_value": 0, "min_period_days": 15, "min_denominator": 1}
START, END = date(2026, 9, 1), date(2026, 9, 30)


def test_missing_data_is_never_zero_defects():
    r = evaluate(ZERO, START, END, 5000, None)
    assert r.result == "insufficient_data" and any("no es cero" in x for x in r.reasons)
    assert evaluate(ZERO, START, END, None, 0).result == "insufficient_data"
    assert allowed_decisions("insufficient_data") == {"insufficient_data"}


def test_rate_against_configurable_threshold():
    assert evaluate(PPM, START, END, 20_000, 2).result == "effective"   # 100 ppm == límite
    r = evaluate(PPM, START, END, 20_000, 3)
    assert r.result == "ineffective" and r.value == pytest.approx(150)
    pct = {**PPM, "threshold_type": "max_rate_percent", "threshold_value": 0.5}
    assert evaluate(pct, START, END, 2000, 10).result == "effective"


def test_criterion_need_not_be_zero_defects_for_30_days():
    lenient = {**ZERO, "threshold_value": 3, "min_period_days": 10}
    assert evaluate(lenient, date(2026, 9, 1), date(2026, 9, 10), 500, 2).result == "effective"


def test_days_without_production_do_not_count():
    r = evaluate(PPM, START, END, 50_000, 0, days_without_production=5)
    assert r.result == "insufficient_data" and "25 día(s)" in r.reasons[0]


def test_minimum_denominator_and_zero_denominator():
    assert evaluate(PPM, START, END, 999, 0).result == "insufficient_data"
    rate = {**PPM, "min_denominator": 1}
    assert evaluate({**rate, "min_denominator": 1}, START, END, 0, 0).result in ("insufficient_data",)


def test_invalid_data_rejected():
    with pytest.raises(InvalidData):
        evaluate(PPM, START, END, 100, 101)  # más defectos que unidades producidas
    with pytest.raises(InvalidData):
        evaluate(PPM, END, START, 100, 1)
    with pytest.raises(InvalidData):
        evaluate(PPM, START, END, 100, 1, days_without_production=40)
    lots = {**PPM, "denominator_type": "lots_received", "min_denominator": 1, "threshold_type": "max_defects", "threshold_value": 50}
    assert evaluate(lots, START, END, 10, 20).result == "effective"  # defectos > lotes es válido con otro denominador


def test_conservative_decisions_only():
    assert allowed_decisions("effective") == {"effective", "ineffective", "insufficient_data"}
    assert "effective" not in allowed_decisions("ineffective")


META = {"title": "Expediente 8D completo", "folio": "8D-P1-2026-0001", "revision": 3, "organization": "Org Ñandú",
        "issued_by": "José Peña", "issued_at": "2026-10-27 10:00 CST", "status_label": "Caso: En proceso",
        "template_version": "8d_full/1.0", "watermark": None}


def test_lesson_pdf_multipage_with_spanish_and_page_numbers():
    long = "Señal de extracción incorrecta; acción: «ajuste» — corrección. " * 120
    lesson = {"title": "Lección: rebaba en carcasa", "folio": "LA-0001", "revision": 2, "status": "published",
              "language_code": "es", "source_folio": "8D-P1-2026-0001", "problem_summary": long,
              "verified_root_cause": "Desgaste", "actions_summary": long, "effectiveness_summary": "Eficaz",
              "before_summary": "Antes", "after_summary": "Después", "documents_updated": ["PFMEA-12 rev C"],
              "training_summary": None, "products": ["Carcasa"], "processes": ["Maquinado"], "plant_codes": ["P1"],
              "situations": [], "created_by_name": "Ana", "created_at": datetime(2026, 10, 1), "reviewed_by_name": "Luis",
              "reviewed_at": datetime(2026, 10, 2), "published_by_name": "Luis", "published_at": datetime(2026, 10, 3),
              "revision_reason": "Actualización", "applications": [{"target_folio": "8D-P2-2026-0004", "revision": 2,
                                                                     "outcome_label": "Aplicada", "evaluation": "Se revisó"}]}
    pdf = pdf_reports.lesson_report(lesson, {**META, "title": "Lección aprendida", "folio": "LA-0001", "revision": 2})
    reader = PdfReader(io.BytesIO(pdf))
    assert len(reader.pages) >= 2
    text = reader.pages[0].extract_text()
    assert "Señal" in text and "Página 1 de" in text and "Org Ñandú" in text
    assert f"Página {len(reader.pages)} de {len(reader.pages)}" in reader.pages[-1].extract_text()


def test_lesson_pdf_before_after_with_many_photos_spans_pages():
    from PIL import Image
    def big(color):
        buf = io.BytesIO()
        Image.new("RGB", (900, 700), color).save(buf, format="PNG")
        return buf.getvalue()
    # Colores distintos: reportlab guarda una sola vez las imágenes idénticas.
    photos = ([("before", f"Antes {i}", big((200, 40 + i * 30, 40))) for i in range(3)]
              + [("after", f"Después {i}", big((40, 160, 60 + i * 30))) for i in range(4)])
    lesson = {"title": "Antes y después", "folio": "LA-0002", "revision": 1, "status": "draft", "language_code": "es",
              "source_folio": "8D-X", "problem_summary": "p", "verified_root_cause": "r", "actions_summary": "a",
              "effectiveness_summary": "e", "before_summary": "Situación previa", "after_summary": "Situación posterior",
              "documents_updated": [], "training_summary": None, "products": [], "processes": [], "plant_codes": [], "situations": [],
              "created_by_name": "A", "created_at": None, "reviewed_by_name": None, "reviewed_at": None, "published_by_name": None,
              "published_at": None, "revision_reason": None, "applications": [], "photo_images": photos}
    reader = PdfReader(io.BytesIO(pdf_reports.lesson_report(lesson, {**META, "title": "Lección aprendida"})))
    text = "".join(pg.extract_text() for pg in reader.pages)
    assert len(reader.pages) >= 2 and "ANTES" in text and "Después 3" in text
    assert sum(len(pg.images) for pg in reader.pages) == 7


def test_excel_escapes_formulas_and_keeps_types():
    content = excel.workbook([("Casos", ["Folio", "Texto", "Fecha", "N"], [["8D-1", "=HYPERLINK(\"x\")", date(2026, 1, 2), 5]])],
                             {**META, "filters": "q=x"})
    wb = load_workbook(io.BytesIO(content))
    assert wb.sheetnames == ["Información", "Casos"]
    row = [c.value for c in wb["Casos"][2]]
    assert row[1].startswith("'=") and row[3] == 5 and row[2].year == 2026
